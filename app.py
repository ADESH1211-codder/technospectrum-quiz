from flask import Flask, request, redirect, session, send_file
import sqlite3, time, random, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

app = Flask(__name__)
app.secret_key = "technospectrum_2k26"

# ---------------- DATABASE ----------------
def db():
    return sqlite3.connect("quiz.db", check_same_thread=False)

con = db()
cur = con.cursor()

cur.execute("""CREATE TABLE IF NOT EXISTS participants(
id INTEGER PRIMARY KEY AUTOINCREMENT,
name TEXT,
email TEXT UNIQUE,
contact TEXT,
score INTEGER DEFAULT 0,
start REAL,
end REAL
)""")

cur.execute("""CREATE TABLE IF NOT EXISTS questions(
id INTEGER PRIMARY KEY AUTOINCREMENT,
q TEXT,a TEXT,b TEXT,c TEXT,d TEXT,correct TEXT
)""")

cur.execute("""CREATE TABLE IF NOT EXISTS settings(
id INTEGER PRIMARY KEY,
timer INTEGER,
quiz_open INTEGER
)""")

if cur.execute("SELECT COUNT(*) FROM settings").fetchone()[0] == 0:
    cur.execute("INSERT INTO settings VALUES(1,300,1)")
    con.commit()

# ---------------- UI TEMPLATE ----------------
def page(body):
    return f"""
<!doctype html>
<html>
<head>
<meta name="viewport" content="width=device-width,initial-scale=1">
<link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.2/dist/css/bootstrap.min.css" rel="stylesheet">
<style>
body{{background:linear-gradient(135deg,#667eea,#764ba2);}}
.card{{border-radius:18px}}
h2{{color:white}}
</style>
</head>
<body>
<div class="container mt-4">
<h2 class="text-center">TECHNOSPECTRUM 2K26</h2>
<div class="card p-4 shadow mt-3">{body}</div>
</div>
</body>
</html>
"""

def quiz_open():
    return cur.execute("SELECT quiz_open FROM settings").fetchone()[0] == 1

def admin_required():
    return session.get("admin")

# ---------------- USER ----------------
@app.route("/", methods=["GET","POST"])
def register():
    if not quiz_open():
        return page("<h4 class='text-center text-danger'>Quiz Closed</h4>")

    if request.method == "POST":
        session["tmp"] = request.form
        return redirect("/otp")

    return page("""
<form method="post">
<input class="form-control mb-2" name="name" placeholder="Full Name" required>
<input class="form-control mb-2" name="email" placeholder="Email" required>
<input class="form-control mb-2" name="contact" placeholder="Contact" required>
<button class="btn btn-primary w-100">Continue</button>
</form>
""")

@app.route("/otp", methods=["GET","POST"])
def otp():
    if request.method == "POST":
        if request.form["otp"] == "1211":
            d = session["tmp"]
            cur.execute(
                "INSERT INTO participants(name,email,contact,start) VALUES(?,?,?,?)",
                (d["name"], d["email"], d["contact"], time.time())
            )
            con.commit()
            session["pid"] = cur.lastrowid
            return redirect("/quiz")
        return page("<h4 class='text-danger text-center'>Wrong OTP</h4>")

    return page("""
<form method="post">
<input class="form-control mb-2" name="otp" placeholder="Enter OTP" required>
<button class="btn btn-success w-100">Verify</button>
</form>
""")

@app.route("/quiz")
def quiz():
    if "pid" not in session:
        return redirect("/")

    qs = cur.execute("SELECT * FROM questions").fetchall()
    random.shuffle(qs)
    timer = cur.execute("SELECT timer FROM settings").fetchone()[0]

    qhtml = ""
    for q in qs:
        qhtml += f"""
<b>{q[1]}</b><br>
<label><input type="radio" name="{q[0]}" value="a"> {q[2]}</label><br>
<label><input type="radio" name="{q[0]}" value="b"> {q[3]}</label><br>
<label><input type="radio" name="{q[0]}" value="c"> {q[4]}</label><br>
<label><input type="radio" name="{q[0]}" value="d"> {q[5]}</label><hr>
"""

    return page(f"""
<div class="alert alert-info text-center">
Time Left: <span id="t">{timer}</span> sec
</div>
<form method="post" action="/submit">
{qhtml}
<button class="btn btn-success w-100">Submit Quiz</button>
</form>

<script>
let t = {timer};
setInterval(function() {{
    document.getElementById("t").innerHTML = t;
    if(t <= 0) {{
        document.forms[0].submit();
    }}
    t--;
}}, 1000);
</script>
""")

@app.route("/submit", methods=["POST"])
def submit():
    score = 0
    for q, a in request.form.items():
        if a == cur.execute("SELECT correct FROM questions WHERE id=?", (q,)).fetchone()[0]:
            score += 1

    cur.execute(
        "UPDATE participants SET score=?, end=? WHERE id=?",
        (score, time.time(), session["pid"])
    )
    con.commit()

    return page("<h4 class='text-center text-success'>You have successfully submitted the quiz</h4>")

# ---------------- ADMIN ----------------
@app.route("/admin", methods=["GET","POST"])
def admin():
    if request.method == "POST":
        if request.form["u"] == "1234" and request.form["p"] == "pass@123":
            session["admin"] = 1

    if not admin_required():
        return page("""
<form method="post">
<input class="form-control mb-2" name="u" placeholder="Username">
<input class="form-control mb-2" type="password" name="p" placeholder="Password">
<button class="btn btn-dark w-100">Login</button>
</form>
""")

    return page("""
<a href="/addq" class="btn btn-primary w-100 mb-2">Add Question</a>
<a href="/settimer" class="btn btn-info w-100 mb-2">Set Timer</a>
<a href="/leaderboard" class="btn btn-secondary w-100 mb-2">Leaderboard</a>
<a href="/graph" class="btn btn-dark w-100 mb-2">Graph</a>
<a href="/export" class="btn btn-success w-100 mb-2">Download Excel</a>
""")

# ---------------- ADMIN FEATURES ----------------
@app.route("/addq", methods=["GET","POST"])
def addq():
    if not admin_required():
        return redirect("/admin")

    if request.method == "POST":
        cur.execute(
            "INSERT INTO questions(q,a,b,c,d,correct) VALUES(?,?,?,?,?,?)",
            (request.form["q"], request.form["a"], request.form["b"],
             request.form["c"], request.form["d"], request.form["correct"])
        )
        con.commit()
        return redirect("/admin")

    return page("""
<form method="post">
<input class="form-control mb-2" name="q" placeholder="Question">
<input class="form-control mb-2" name="a" placeholder="Option A">
<input class="form-control mb-2" name="b" placeholder="Option B">
<input class="form-control mb-2" name="c" placeholder="Option C">
<input class="form-control mb-2" name="d" placeholder="Option D">
<select class="form-control mb-2" name="correct">
<option>a</option><option>b</option><option>c</option><option>d</option>
</select>
<button class="btn btn-success w-100">Add Question</button>
</form>
""")

@app.route("/leaderboard")
def leaderboard():
    if not admin_required():
        return redirect("/admin")

    rows = cur.execute("SELECT name,score FROM participants ORDER BY score DESC").fetchall()
    t = ""
    for i, r in enumerate(rows):
        t += f"<tr><td>{i+1}</td><td>{r[0]}</td><td>{r[1]}</td></tr>"

    return page(f"""
<table class="table">
<tr><th>Rank</th><th>Name</th><th>Score</th></tr>
{t}
</table>
""")

@app.route("/graph")
def graph():
    if not admin_required():
        return redirect("/admin")

    rows = cur.execute("SELECT name,score FROM participants").fetchall()
    names = [r[0] for r in rows]
    scores = [r[1] for r in rows]

    return page(f"""
<canvas id="chart"></canvas>
<script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
<script>
new Chart(document.getElementById("chart"), {{
    type: "bar",
    data: {{
        labels: {names},
        datasets: [{{
            label: "Scores",
            data: {scores},
            backgroundColor: "#6f42c1"
        }}]
    }}
}});
</script>
""")

@app.route("/export")
def export():
    if not admin_required():
        return redirect("/admin")

    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:E1")
    ws["A1"] = "TECHNOSPECTRUM 2K26 – QUIZ RESULTS"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.append(["Rank","Name","Email","Contact","Score"])
    for c in ws[2]:
        c.font = Font(bold=True)

    rows = cur.execute(
        "SELECT name,email,contact,score FROM participants ORDER BY score DESC"
    ).fetchall()

    for i, r in enumerate(rows):
        ws.append([i+1, r[0], r[1], r[2], r[3]])

    for col in "ABCDE":
        ws.column_dimensions[col].width = 22

    wb.save("Technospectrum_2K26.xlsx")
    return send_file("Technospectrum_2K26.xlsx", as_attachment=True)

@app.route("/settimer", methods=["GET","POST"])
def settimer():
    if not admin_required():
        return redirect("/admin")

    if request.method == "POST":
        total = int(request.form["m"]) * 60 + int(request.form["s"])
        cur.execute("UPDATE settings SET timer=?", (total,))
        con.commit()
        return redirect("/admin")

    return page("""
<form method="post">
<input class="form-control mb-2" name="m" placeholder="Minutes">
<input class="form-control mb-2" name="s" placeholder="Seconds">
<button class="btn btn-info w-100">Set Timer</button>
</form>
""")

app.run(host="0.0.0.0", port=5000)
